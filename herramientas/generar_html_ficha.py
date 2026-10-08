"""Genera frontend/static/ficha_matricula.html: la ficha oficial como tablas
HTML reales (mismas celdas, colores y fusiones que el Word), con un <input>
dentro de cada celda de dato. Sale del Word ya limpio (sin datos de nadie),
así las etiquetas quedan completas y exactas. Uso: python herramientas/generar_html_ficha.py"""
import html, io, sys
from pathlib import Path

RAIZ = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(RAIZ / "backend"))
from docx import Document
from docx.oxml.ns import qn
from app import ficha_matricula as fm

# Word "vacío": ningún dato de persona, solo etiquetas y valores fijos.
vacio = fm.generar_ficha_matricula(
    documento="", tipo_documento=None, lugar_expedicion_documento=None, primer_nombre="", segundo_nombre=None,
    primer_apellido="", segundo_apellido=None, fecha_nacimiento=None, genero=None, grado=0, curso=None,
    campos={"grado_2026": "", "dilig_dia": "", "dilig_mes": "", "dilig_anio": ""},
)
doc = Document(io.BytesIO(vacio))
CAMPO = dict(fm.MAPA_CAMPOS); CAMPO["numero_documento"] = (0, 7, 1)
por_pos = {v: k for k, v in CAMPO.items()}


def info_celda(tc):
    fill = None
    shd = tc.tcPr.find(qn("w:shd")) if tc.tcPr is not None else None
    if shd is not None and shd.get(qn("w:fill")) not in (None, "auto"):
        fill = "#" + shd.get(qn("w:fill"))
    txt, size, bold, jc = [], None, False, None
    for p in tc.iter(qn("w:p")):
        t = "".join(x.text or "" for x in p.iter(qn("w:t"))).strip()
        if t:
            txt.append(t)
        for r in p.iter(qn("w:r")):
            if size is None and r.find(qn("w:rPr")) is not None:
                sz = r.find(qn("w:rPr")).find(qn("w:sz"))
                if sz is not None: size = int(sz.get(qn("w:val"))) / 2
            if r.find(qn("w:rPr")) is not None and r.find(qn("w:rPr")).find(qn("w:b")) is not None: bold = True
        ppr = p.find(qn("w:pPr"))
        if ppr is not None and ppr.find(qn("w:jc")) is not None and jc is None: jc = ppr.find(qn("w:jc")).get(qn("w:val"))
    return fill, txt, size or 8, bold, jc


def tabla_html(ti, tabla):
    grid = [int(g.w) for g in tabla._tbl.tblGrid.gridCol_lst]; total = sum(grid) or 1
    filas = list(tabla.rows)
    inicio = []  # por fila: [(tc, col_inicio, span, vmerge)]
    for fila in filas:
        c, lst = 0, []
        for tc in fila._tr.tc_lst:
            span = tc.grid_span or 1
            lst.append((tc, c, span, tc.vMerge)); c += span
        inicio.append(lst)
    # Una celda fusionada en vertical aparece en varias filas: el nombre del
    # campo se asigna a la celda física (tc), sea cual sea la fila que lo cite.
    nombre_de = {}
    for r, fila in enumerate(filas):
        for i, u in enumerate(fm._celdas_unicas(fila)):
            n = por_pos.get((ti, r, i))
            if n: nombre_de.setdefault(id(u._tc), n)
    out = ['<table class="fh-t"><colgroup>' + "".join(f'<col style="width:{w / total * 100:.3f}%">' for w in grid) + "</colgroup>"]
    for r, fila in enumerate(filas):
        out.append("<tr>")
        for tc, c0, span, vm in inicio[r]:
            if vm == "continue":
                continue
            rs = 1
            if vm == "restart":
                k = r + 1
                while k < len(filas) and any(c == c0 and v == "continue" for _, c, _, v in inicio[k]): rs += 1; k += 1
            fill, txt, size, bold, jc = info_celda(tc)
            nombre = nombre_de.get(id(tc))
            estilo = f"font-size:{size / 8:.2f}em;" + (f"background:{fill};" if fill else "")
            if jc in ("center", "right"): estilo += f"text-align:{jc};"
            attrs = (f' colspan="{span}"' if span > 1 else "") + (f' rowspan="{rs}"' if rs > 1 else "")
            if nombre:
                out.append(f'<td{attrs} class="fh-v" style="{estilo}"><input name="{nombre}" autocomplete="off"></td>')
            else:
                out.append(f'<td{attrs} style="{estilo}{"font-weight:700;" if bold else ""}">' + "<br>".join(html.escape(t) for t in txt) + "</td>")
        out.append("</tr>")
    return "".join(out) + "</table>"


motivos = ""
try:
    origen = Document(r"D:\ietaba\ARCHIVO INGENIERO\1 FICHA MATRICULA INST. EDUCA. TEC. AGRO. BILINGUE - IETABA.docx")
    motivos = origen.paragraphs[5].text.strip()
except Exception:
    pass

css = """<style>
.fh{--z:1.3;width:calc(680px*var(--z));max-width:none;background:#fff;color:#111;font-family:Cambria,Georgia,serif;font-size:calc(9px*var(--z));padding:calc(10px*var(--z))}
.fh-enc{display:flex;align-items:center;gap:8px;text-align:center;font-family:'Arial Narrow',Arial,sans-serif;font-weight:700;line-height:1.25}
.fh-enc img{width:calc(48px*var(--z));height:calc(48px*var(--z));object-fit:contain}
.fh-enc div{flex:1}.fh-enc .g{font-size:1.25em}
.fh-t{width:100%;border-collapse:collapse;table-layout:fixed;margin:calc(6px*var(--z)) 0}
.fh-t td{border:.75px solid #000;padding:1px 3px;vertical-align:middle;overflow:hidden;line-height:1.15;height:1.9em}
.fh-t td.fh-v{padding:0;background:#fff}
.fh-t input{display:block;box-sizing:border-box;width:100%;height:100%;min-height:1.9em;border:0;outline:0;background:transparent;font:700 1em Cambria,Georgia,serif;padding:0 3px;color:#0b2a7a}
.fh-t input:hover{background:rgba(255,255,170,.5)}.fh-t input:focus{background:rgba(255,255,140,.9);outline:1px dashed #777}
.fh-pie{text-align:justify;font-size:.95em;line-height:1.25;margin-top:6px}
@media print{@page{margin:8mm}.fh{width:100%!important;font-size:8.5pt!important;padding:0}.fh-t input{color:#000}}
</style>"""
enc = ('<div class="fh-enc"><img src="/static/logos/logo2.jpeg" alt=""><div><div class="g">INSTITUCION EDUCATIVA INDIGENA TECNICA AGROAMBIENTAL BILINGÜE AWA – IETABA</div>'
       '<div class="g">SEDE # 01 INSTITUCIÓN IETABA - CODIGO DANE 252079002045</div><div>MUNICIPIO DE BARBACOAS - 2027</div>'
       '<div>FORMATO PARA MATRICULA DE ESTUDIANTES AÑO ESCOLAR 2027</div></div></div>')
cuerpo = "".join(tabla_html(i, t) for i, t in enumerate(doc.tables[:5]))
(RAIZ / "frontend/static/ficha_matricula.html").write_text(
    css + f'<div class="fh" id="ficha-hoja">{enc}{cuerpo}<div class="fh-pie">{html.escape(motivos)}</div></div>', encoding="utf-8")
print("ok", cuerpo.count("<input"), "casillas")
