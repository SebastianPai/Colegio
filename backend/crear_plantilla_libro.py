"""Crea plantillas/libro_matricula.docx a partir del Word oficial del colegio
("LIBRO DE MATRICULA IETABA 2026..docx"), SIN tocar el original.

La copia queda con UNA sola hoja (el primer registro) y sin ningún dato real:
cada dato del estudiante se cambia por un marcador «campo» que
app/libro_matricula.py reemplaza al generar. Membrete (encabezado), pie de
página, márgenes, letra y textos fijos quedan exactamente como en el original.

    ./venv/Scripts/python.exe crear_plantilla_libro.py "D:\\...\\LIBRO DE MATRICULA IETABA 2026..docx"
"""

import shutil
import sys
from pathlib import Path

from docx import Document

DESTINO = Path(__file__).resolve().parent / "plantillas" / "libro_matricula.docx"
PARRAFOS_POR_HOJA = 32  # cada registro del libro ocupa 32 párrafos


def _poner(p, textos: dict[int, str]) -> None:
    """Cambia el texto de las corridas indicadas (las demás no se tocan)."""
    for i, texto in textos.items():
        p.runs[i].text = texto


def main(original: str) -> None:
    shutil.copyfile(original, DESTINO)
    doc = Document(DESTINO)
    cuerpo = doc.element.body
    parrafos = doc.paragraphs
    # Solo la primera hoja: se quitan los demás registros (se deja sectPr).
    for p in parrafos[PARRAFOS_POR_HOJA:]:
        cuerpo.remove(p._p)
    p = doc.paragraphs
    marca = lambda c: f"_«{c}»_"  # noqa: E731  — "SI___" en blanco, "SI_X_" marcado

    _poner(p[2], {1: "«anio»", 2: ""})
    _poner(p[5], {3: "«nombre»", 4: "", 5: ""})
    _poner(p[6], {2: "«nacimiento»"})
    _poner(p[7], {1: "«documento»", 8: "«expedido»"})
    _poner(p[8], {2: "«direccion»"})
    _poner(p[9], {1: "«padres»", 2: "", 3: "", 4: "", 5: "", 6: ""})
    _poner(p[10], {2: "«grado»", 5: "__«anio»", 6: ""})
    _poner(p[11], {
        0: f"SITUACION ACTUAL DEL ESTUDIANTE: DESPLAZADO    SI{marca('desplaz_si')} NO ",
        1: marca("desplaz_no"), 2: "", 3: "", 4: "",
        7: f"\tHUERFANO    SI{marca('huerfano_si')} NO ", 8: marca("huerfano_no"), 9: "",
        11: f"NECESITA HOSPEDAJE EN EL ALBERGUE SI {marca('albergue_si')} NO",
        13: marca("albergue_no"), 14: "", 15: "",
        16: f"DISCAPACIDAD:   COGNITIVA: {marca('disc_cognitiva')}        SORDOMUDO {marca('disc_sordera')}"
            f"          DIFICULTAD PARA MOVERSE {marca('disc_fisica')}         CEGUERA {marca('disc_ceguera')}",
    })
    _poner(p[12], {2: "«telefono»"})
    _poner(p[13], {0: "ESTUDIANTE ANTIGUO ____«est_antiguo»", 5: "\tESTUDIANTE NUEVO ___", 6: "«est_nuevo»"})
    _poner(p[14], {0: "PROVIENE DEL C. EDUCATIVO: ", 1: "", 2: "", 3: "«procedencia»", 4: "", 5: "", 6: "", 7: ""})
    _poner(p[15], {2: "«dilig_dia»", 6: "«dilig_mes»", 8: "___«dilig_anio»", 9: ""})
    _poner(p[19], {0: "«firma_nombre_est»________"})
    _poner(p[23], {0: "«acudiente» "})
    _poner(p[27], {0: "«responsable»"})
    doc.save(DESTINO)
    print(f"Plantilla creada: {DESTINO}")


if __name__ == "__main__":
    main(sys.argv[1])
