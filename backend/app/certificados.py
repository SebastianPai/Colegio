"""Genera constancias de estudio a partir de la plantilla oficial del colegio
(plantillas/constancia_estudio.docx), sin tocar su formato: solo se
reemplaza el texto de los "runs" (fragmentos de texto con estilo) que
contienen los datos variables, dejando el resto del documento intacto
(membrete, logos, firma del rector).
"""

import datetime as dt
import io
from pathlib import Path

from docx import Document

from .models import GRADOS_NOMBRES

PLANTILLA = Path(__file__).resolve().parent.parent / "plantillas" / "constancia_estudio.docx"


def convertir_a_pdf(ruta_docx: Path, ruta_pdf: Path) -> bool:
    """Convierte usando Word (instalado en este PC). Si Word no está
    disponible, devuelve False y el sistema sigue funcionando solo con
    Word (.docx) — la conversión a PDF es un extra, no un requisito."""
    # Word por COM, en una instancia PROPIA (DispatchEx) que se cierra al
    # terminar. Antes (docx2pdf) se enganchaba al Word que la persona tuviera
    # abierto en el PC: si ese Word estaba ocupado o con un aviso abierto, la
    # conversión fallaba (así quedaban muchas constancias solo en Word) y el
    # archivo se quedaba abierto y bloqueado en ese Word. Además los endpoints
    # corren en hilos del servidor, donde COM no está iniciado: se inicia aquí.
    try:
        import pythoncom
        import win32com.client
    except ImportError:
        return False
    pythoncom.CoInitialize()
    word = documento = None
    try:
        word = win32com.client.DispatchEx("Word.Application")
        word.Visible = False
        word.DisplayAlerts = 0
        documento = word.Documents.Open(str(ruta_docx.resolve()), ReadOnly=True, AddToRecentFiles=False)
        documento.ExportAsFixedFormat(str(ruta_pdf.resolve()), 17)  # 17 = wdExportFormatPDF
        return ruta_pdf.exists()
    except Exception:
        return False
    finally:
        try:
            if documento is not None:
                documento.Close(False)
        except Exception:
            pass
        try:
            if word is not None:
                word.Quit(False)
        except Exception:
            pass
        pythoncom.CoUninitialize()

NIVELES_POR_GRADO = {
    0: "preescolar",
    1: "básica primaria",
    2: "básica primaria",
    3: "básica primaria",
    4: "básica primaria",
    5: "básica primaria",
    6: "básica secundaria",
    7: "básica secundaria",
    8: "básica secundaria",
    9: "básica secundaria",
    10: "media",
    11: "media",
}

DIAS_TEXTO = {
    1: "uno", 2: "dos", 3: "tres", 4: "cuatro", 5: "cinco", 6: "seis", 7: "siete",
    8: "ocho", 9: "nueve", 10: "diez", 11: "once", 12: "doce", 13: "trece",
    14: "catorce", 15: "quince", 16: "dieciséis", 17: "diecisiete", 18: "dieciocho",
    19: "diecinueve", 20: "veinte", 21: "veintiuno", 22: "veintidós", 23: "veintitrés",
    24: "veinticuatro", 25: "veinticinco", 26: "veintiséis", 27: "veintisiete",
    28: "veintiocho", 29: "veintinueve", 30: "treinta", 31: "treinta y uno",
}

MESES_TEXTO = {
    1: "enero", 2: "febrero", 3: "marzo", 4: "abril", 5: "mayo", 6: "junio",
    7: "julio", 8: "agosto", 9: "septiembre", 10: "octubre", 11: "noviembre",
    12: "diciembre",
}

TIPOS_DOCUMENTO_TEXTO = {
    "C.C.": "Cédula de ciudadanía",
    "T.I.": "Tarjeta de identidad",
    "R.C.": "Registro civil de nacimiento",
}

PROPOSITOS = [
    "Trámite de libreta militar",
    "Trámite ante el ICBF",
    "Trámite bancario",
    "Solicitud de subsidio o beneficio del gobierno",
    "Constancia de estudio general",
]


def _texto_tipo_documento(tipo_documento: str | None) -> str:
    if not tipo_documento:
        return "Cédula de ciudadanía"
    return TIPOS_DOCUMENTO_TEXTO.get(tipo_documento, tipo_documento)


def _fusionar_runs(paragraph, inicio: int, fin: int, texto_nuevo: str) -> None:
    """Pone `texto_nuevo` en el run `inicio` (conservando su formato) y
    vacía los runs desde `inicio+1` hasta `fin` (inclusive)."""
    runs = paragraph.runs
    runs[inicio].text = texto_nuevo
    for i in range(inicio + 1, fin + 1):
        runs[i].text = ""


def generar_constancia(
    *,
    nombre_completo: str,
    tipo_documento: str | None,
    documento: str,
    lugar_expedicion: str,
    grado: int,
    jornada: str,
    anio_escolar: int,
    fecha_expedicion: dt.date,
    proposito: str,
) -> bytes:
    doc = Document(PLANTILLA)
    parrafos = doc.paragraphs

    nombre_grado = GRADOS_NOMBRES.get(grado, str(grado)).lower()
    nivel = NIVELES_POR_GRADO.get(grado, "básica")

    # Párrafo del cuerpo: "Que el estudiante: NOMBRE, identificado(a) con
    # TIPO_DOC No. DOCUMENTO de LUGAR se encuentra matriculado... en el
    # Grado PALABRA (NUMERO) de la Educación NIVEL, jornada JORNADA,
    # Calendario A, año escolar AÑO."
    p_cuerpo = parrafos[6]
    _fusionar_runs(p_cuerpo, 1, 1, nombre_completo)
    _fusionar_runs(p_cuerpo, 3, 3, _texto_tipo_documento(tipo_documento))
    _fusionar_runs(p_cuerpo, 7, 7, documento)
    _fusionar_runs(p_cuerpo, 9, 11, f"de {lugar_expedicion}")
    _fusionar_runs(p_cuerpo, 26, 26, nombre_grado)
    _fusionar_runs(p_cuerpo, 28, 28, str(grado))
    _fusionar_runs(p_cuerpo, 30, 30, nivel)
    _fusionar_runs(p_cuerpo, 34, 34, f" {jornada}")
    _fusionar_runs(p_cuerpo, 35, 36, f", Calendario A, año escolar {anio_escolar}")

    # Párrafo de la fecha: "...a los DIA (NUMERO) días del mes de MES del
    # año AÑO."
    p_fecha = parrafos[8]
    _fusionar_runs(p_fecha, 4, 4, DIAS_TEXTO.get(fecha_expedicion.day, str(fecha_expedicion.day)))
    _fusionar_runs(p_fecha, 6, 8, f"({fecha_expedicion.day}) día")
    _fusionar_runs(p_fecha, 11, 11, MESES_TEXTO.get(fecha_expedicion.month, ""))
    _fusionar_runs(p_fecha, 12, 13, f" del año {fecha_expedicion.year}")

    # Párrafo "Válido para ..."
    p_valido = parrafos[11]
    _fusionar_runs(p_valido, 0, 3, f"Válido para {proposito}.")

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


# --- Paz y salvo de servicio social (grado 11) -------------------------------
# Plantilla: plantillas/paz_y_salvo_servicio_social.docx, copia del Word
# oficial "PAZ Y SALVO 2026" (certifica las 80 horas de servicio social
# obligatorio, requisito para el título de bachiller). Igual que la
# constancia: solo se cambian los runs con datos; membrete, logos, texto
# legal y firma del rector quedan intactos.

PLANTILLA_PAZ_Y_SALVO = PLANTILLA.parent / "paz_y_salvo_servicio_social.docx"
PROPOSITO_PAZ_Y_SALVO = "Paz y salvo — servicio social (80 horas)"

# Índices dentro del cuerpo de la plantilla (ver la estructura de runs en
# generar_paz_y_salvo): párrafo del "Que el estudiante..." y el de la fecha.
_P_CUERPO_PYS, _P_FECHA_PYS = 5, 7
_ULTIMO_PARRAFO_PYS = 15  # "Rector IETABA": lo de abajo son renglones vacíos

TIPOS_DOCUMENTO_CORTO = {"C.C.": "C.C", "T.I.": "T.I", "R.C.": "R.C"}


def _documento_con_puntos(documento: str) -> str:
    """'1087185389' -> '1.087.185.389' (como en el formato oficial). Lo que
    no sea solo dígitos se deja tal cual."""
    limpio = (documento or "").strip()
    if not limpio.isdigit():
        return limpio
    return f"{int(limpio):,}".replace(",", ".")


def _llenar_paz_y_salvo(parrafos, est: dict, fecha: dt.date) -> None:
    """Runs de la plantilla:
    cuerpo: [0]'Que el estudiante ' [1]NOMBRE(negrita) [2]' identificado con C.C No: '
            [3]'1' [4]'.087.185.389 ' [5]'expedida en Tumaco,' [6]'cumplió...'
    fecha:  [0]'Para constancia se firma a los' [1]' 22 días d' [2]'el mes de octubre del año 2026' [3]'.'"""
    mujer = (est.get("genero") or "").upper() == "F"
    tipo = TIPOS_DOCUMENTO_CORTO.get(est.get("tipo_documento") or "", est.get("tipo_documento") or "documento")
    lugar = (est.get("lugar_expedicion") or "").strip()
    p = parrafos[_P_CUERPO_PYS]
    _fusionar_runs(p, 0, 0, "Que la estudiante " if mujer else "Que el estudiante ")
    _fusionar_runs(p, 1, 1, est["nombre_completo"].upper())
    _fusionar_runs(p, 2, 2, f" {'identificada' if mujer else 'identificado'} con {tipo} No: ")
    _fusionar_runs(p, 3, 4, f"{_documento_con_puntos(est['documento'])} " if lugar else f"{_documento_con_puntos(est['documento'])},")
    _fusionar_runs(p, 5, 5, f"expedida en {lugar}," if lugar else "")
    p_fecha = parrafos[_P_FECHA_PYS]
    if fecha.day == 1:  # "a los 1 día" no se dice: "el primer día del mes..."
        _fusionar_runs(p_fecha, 0, 0, "Para constancia se firma el")
        _fusionar_runs(p_fecha, 1, 1, " primer día d")
    else:
        _fusionar_runs(p_fecha, 1, 1, f" {fecha.day} días d")
    _fusionar_runs(p_fecha, 2, 2, f"el mes de {MESES_TEXTO[fecha.month]} del año {fecha.year}")


def generar_paz_y_salvo(estudiantes: list[dict], fecha_expedicion: dt.date) -> bytes:
    """Un solo Word con una página por estudiante (el mismo membrete y pie de
    página en todas). `estudiantes`: dicts con nombre_completo, documento,
    tipo_documento, lugar_expedicion y genero."""
    import copy

    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.text.paragraph import Paragraph

    if not estudiantes:
        raise ValueError("No hay estudiantes para el paz y salvo.")
    doc = Document(PLANTILLA_PAZ_Y_SALVO)
    cuerpo = doc.element.body
    seccion = cuerpo.find(qn("w:sectPr"))
    pagina = [e for e in cuerpo if e is not seccion]
    parrafos_pagina = [e for e in pagina if e.tag == qn("w:p")]
    # Los renglones vacíos del final empujaban cada página a una hoja de más.
    pagina = pagina[: pagina.index(parrafos_pagina[_ULTIMO_PARRAFO_PYS]) + 1]
    for e in list(cuerpo):
        if e is not seccion:
            cuerpo.remove(e)

    for i, est in enumerate(estudiantes):
        elementos = [copy.deepcopy(e) for e in pagina]
        if i > 0:  # cada estudiante empieza en hoja nueva
            primero = next(e for e in elementos if e.tag == qn("w:p"))
            ppr = primero.get_or_add_pPr()
            # El esquema de Word exige este orden: pStyle, keepNext, keepLines, pageBreakBefore...
            antes = [c for c in ppr if c.tag in (qn("w:pStyle"), qn("w:keepNext"), qn("w:keepLines"))]
            ppr.insert(ppr.index(antes[-1]) + 1 if antes else 0, OxmlElement("w:pageBreakBefore"))
        for e in elementos:
            seccion.addprevious(e)
        parrafos = [Paragraph(e, doc._body) for e in elementos if e.tag == qn("w:p")]
        _llenar_paz_y_salvo(parrafos, est, fecha_expedicion)

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()
