"""Reporte de horario en Word (.docx) con el membrete oficial.

La secretaría arma una lista de "bloques" (día + grado + horas) y este módulo
devuelve un .docx: por cada bloque, una tabla por cada división activa del
grado en la Sede #01 (11-1, 11-2, …), con una fila por hora pedida
(Hora · Docente · Materia). Si a esa hora el grupo no tiene clase en el
horario oficial, la fila igual aparece con "Sin clase".

No inventa datos: todo sale de BloqueHorario (el Word oficial ya importado) y
de las divisiones que existen entre los estudiantes ACTIVO de la sede.
"""

import datetime as dt
import io
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt, RGBColor
from sqlalchemy import select
from sqlalchemy.orm import Session

from . import horario, models

LOGOS = Path(__file__).resolve().parent.parent.parent / "frontend" / "static" / "logos"

DIA_TOKENS = {"lu": 1, "ma": 2, "mi": 3, "ju": 4, "vi": 5}

MEMBRETE_LINEAS = [
    ('UNIDAD INDIGENA DEL PUEBLO AWA "UNIPA"', "titulo"),
    ('INSTITUCION EDUCATIVA INDIGENA TECNICA AGROAMBIENTAL BILINGÜE AWA "IETABA"', "titulo"),
    ("Licencia de Funcionamiento No. 398 del 28 de abril del 2004", "chico"),
    ("Emanada de la Secretaría Departamental de Educación y Cultura de Nariño", "chico"),
    ("Código DANE 25207900204501 - NIT. 900000095-4", "chico"),
    ("Ambiente – Cultura – Ciencia", "lema"),
]


class BloqueReporte:
    """Un renglón de la solicitud: 'jueves, grado 11, horas 2-3-4'."""

    def __init__(self, dia: int, es_hoy: bool, grado: int, horas: list[int]):
        self.dia = dia
        self.es_hoy = es_hoy
        self.grado = grado
        self.horas = horas


def parsear_bloques(tokens: list[str], fecha_hoy: dt.date) -> list[BloqueReporte]:
    """Cada token es 'DIA-GRADO-H-H-…' — p. ej. 'ju-11-2-3-4' o 'hoy-8-4-5-6'."""
    bloques: list[BloqueReporte] = []
    for crudo in tokens:
        partes = [p.strip().lower() for p in crudo.split("-") if p.strip()]
        if len(partes) < 3:
            raise ValueError(f"Bloque mal formado: «{crudo}» (se espera DIA-GRADO-HORAS)")
        dia_tok, grado_tok, *horas_tok = partes

        if dia_tok == "hoy":
            dia = fecha_hoy.weekday() + 1
            if dia > 5:
                raise ValueError("No se puede usar «hoy»: hoy es fin de semana, elige un día concreto.")
            es_hoy = True
        elif dia_tok in DIA_TOKENS:
            dia = DIA_TOKENS[dia_tok]
            es_hoy = False
        else:
            raise ValueError(f"Día no reconocido en «{crudo}» (usa lu, ma, mi, ju, vi o hoy)")

        try:
            grado = int(grado_tok)
        except ValueError:
            raise ValueError(f"Grado inválido en «{crudo}»")
        if grado not in models.GRADOS_NOMBRES:
            raise ValueError(f"Grado fuera de rango (0–11) en «{crudo}»")

        horas: list[int] = []
        for h in horas_tok:
            try:
                n = int(h)
            except ValueError:
                raise ValueError(f"Hora inválida «{h}» en «{crudo}»")
            if not 1 <= n <= len(horario.FRANJAS):
                raise ValueError(f"Hora fuera de rango (1–{len(horario.FRANJAS)}) en «{crudo}»")
            if n not in horas:
                horas.append(n)
        if not horas:
            raise ValueError(f"El bloque «{crudo}» no trae ninguna hora")

        bloques.append(BloqueReporte(dia, es_hoy, grado, sorted(horas)))
    return bloques


def _divisiones_activas(db: Session, grado: int) -> list[int | None]:
    filas = db.scalars(
        select(models.Estudiante.curso)
        .where(
            models.Estudiante.sede_id == horario.SEDE_HORARIO_ID,
            models.Estudiante.grado == grado,
            models.Estudiante.estado == "ACTIVO",
        )
        .distinct()
    ).all()
    cursos = sorted({c for c in filas if c is not None})
    return cursos or [None]  # sin divisiones registradas: el grado va completo


def _clase_en(
    db: Session, anio: int, grado: int, curso: int | None, dia: int, hora: int
) -> tuple[str, str]:
    """(docente, materia) para ese grupo/día/hora. Prefiere el bloque de la
    división exacta; si no hay, usa el del grado completo (curso NULL). Si dos
    docentes comparten la franja (conflicto sin resolver), los une con « / »."""
    bloques = db.scalars(
        select(models.BloqueHorario).where(
            models.BloqueHorario.anio_lectivo == anio,
            models.BloqueHorario.grado == grado,
            models.BloqueHorario.dia == dia,
            models.BloqueHorario.hora == hora,
            models.BloqueHorario.estado == "activo",
        )
    ).all()
    elegidos = [b for b in bloques if b.curso == curso]
    if not elegidos and curso is not None:
        elegidos = [b for b in bloques if b.curso is None]
    if not elegidos:
        return "—", "Sin clase"
    docentes = " / ".join(dict.fromkeys(b.docente.nombre_completo for b in elegidos if b.docente))
    materias = " / ".join(
        dict.fromkeys((b.area.nombre if b.area else b.area_label) for b in elegidos)
    )
    return docentes or "—", materias or "—"


# --- Armado del .docx -----------------------------------------------------


def _membrete(doc: Document) -> None:
    tabla = doc.add_table(rows=1, cols=3)
    tabla.alignment = WD_TABLE_ALIGNMENT.CENTER
    tabla.autofit = False
    izq, centro, der = tabla.rows[0].cells
    izq.width = Cm(3)
    centro.width = Cm(11)
    der.width = Cm(3)

    for celda, logo in ((izq, "logo1.jpg"), (der, "logo2.jpeg")):
        ruta = LOGOS / logo
        p = celda.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if ruta.exists():
            p.add_run().add_picture(str(ruta), width=Cm(2.4))

    centro.paragraphs[0].text = ""
    for i, (texto, tipo) in enumerate(MEMBRETE_LINEAS):
        p = centro.paragraphs[0] if i == 0 else centro.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(texto)
        if tipo == "titulo":
            run.bold = True
            run.font.size = Pt(9)
        elif tipo == "lema":
            run.italic = True
            run.font.size = Pt(8)
        else:
            run.font.size = Pt(7)

    doc.add_paragraph()


def _titulo_bloque(b: BloqueReporte, curso: int | None, fecha_hoy: dt.date) -> str:
    grupo = horario.etiqueta_grupo(b.grado, curso)
    nombre_grado = models.GRADOS_NOMBRES.get(b.grado, str(b.grado))
    dia_nombre = horario.DIAS_NOMBRES.get(b.dia, str(b.dia))
    sufijo = f" (hoy, {fecha_hoy.strftime('%d/%m/%Y')})" if b.es_hoy else ""
    return f"{nombre_grado} {grupo} — {dia_nombre}{sufijo}"


def _tabla_grupo(
    doc: Document, db: Session, anio: int, b: BloqueReporte, curso: int | None, fecha_hoy: dt.date
) -> None:
    sub = doc.add_paragraph()
    sub.paragraph_format.space_before = Pt(6)
    r = sub.add_run(_titulo_bloque(b, curso, fecha_hoy))
    r.bold = True
    r.font.size = Pt(11)

    tabla = doc.add_table(rows=1, cols=3)
    tabla.style = "Table Grid"
    tabla.alignment = WD_TABLE_ALIGNMENT.CENTER
    for celda, texto, ancho in (
        (tabla.rows[0].cells[0], "Hora", Cm(3.5)),
        (tabla.rows[0].cells[1], "Docente", Cm(7.5)),
        (tabla.rows[0].cells[2], "Materia", Cm(6)),
    ):
        celda.width = ancho
        run = celda.paragraphs[0].add_run(texto)
        run.bold = True

    franjas = {f["hora"]: f for f in horario.FRANJAS}
    for hora in b.horas:
        f = franjas.get(hora, {"inicio": "", "fin": ""})
        docente, materia = _clase_en(db, anio, b.grado, curso, b.dia, hora)
        fila = tabla.add_row().cells
        fila[0].text = f"{hora}ª  {f['inicio']}–{f['fin']}"
        fila[1].text = docente
        fila[2].text = materia
        if materia == "Sin clase":
            for celda in fila:
                for p in celda.paragraphs:
                    for run in p.runs:
                        run.font.color.rgb = RGBColor(0x88, 0x88, 0x88)


def generar_reporte_horario_docx(
    db: Session, tokens: list[str], anio_lectivo: int, fecha_hoy: dt.date
) -> bytes:
    bloques = parsear_bloques(tokens, fecha_hoy)
    if not bloques:
        raise ValueError("No se recibió ningún bloque para el reporte.")

    doc = Document()
    _membrete(doc)

    tit = doc.add_paragraph()
    tit.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = tit.add_run("REPORTE DE HORARIO")
    run.bold = True
    run.font.size = Pt(13)
    pie = doc.add_paragraph()
    pie.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pie.add_run(
        f"Sede #01 · Año lectivo {anio_lectivo} · Generado el {fecha_hoy.strftime('%d/%m/%Y')}"
    ).font.size = Pt(9)

    for b in bloques:
        for curso in _divisiones_activas(db, b.grado):
            _tabla_grupo(doc, db, anio_lectivo, b, curso, fecha_hoy)
        doc.add_paragraph()

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()
